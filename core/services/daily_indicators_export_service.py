"""Exportacao CSV/XLSX/ZIP do relatorio Indicadores diarios (REL-13-E).

Fonte da verdade: docs/superpowers/specs/2026-09-22-reports-hub-design.md,
secoes REL-07 (contrato comum) e REL-13-C (contrato especifico, decisoes
A-D de 2026-09-30). Reaproveita o motor de REL-13-Q
(core/services/daily_indicators_report_service): mesmo calculo
(`build_indicator_for_day`), mesma granularidade por dia calendario, mesma
projecao de linha (`REPORT_ROW_KEYS`) — nenhuma formula redefinida, nenhuma
escrita de negocio, model ou migration.

Exportador PROPRIO de Indicadores diarios, mesmo padrao arquitetural ja
usado por Selos WhatsApp/Inventario/Alocacoes/Pendencias/Historico da linha
(streaming + limpeza de temporarios + divisao em ZIP + protecao de
formula), com seu proprio COLUMNS/slug/row-mapping, deliberadamente
duplicado em vez de generalizado (mesma decisao ja registrada em
REL-09-C/REL-10-E: cada relatorio com seu proprio exportador).

Diferenca estrutural dos outros cinco: a granularidade e por DIA
CALENDARIO, nao por registro-fonte — `iter_daily_indicators_report_rows`
(daily_indicators_report_service) itera um dia por vez, sem paginacao,
reexecutando `build_indicator_for_day` (~8-10 consultas/dia) para TODOS os
dias do periodo, nao so os da pagina atual (diferente da tela, REL-13-Q).
O limiar de particionamento (`EXPORT_SPLIT_THRESHOLD`) conta DIAS
exportados, nao registros-fonte — 100.000 registros-fonte nao equivalem a
100.000 linhas exportadas aqui (a granularidade e estrutural, uma linha por
dia do periodo, nunca por entidade).

Otimizacao aplicada (auditoria pos-REL-13-Q, "Requisito de investigacao
explicito para REL-13-E"): a exportacao processa o periodo INTEIRO sem o
teto de pagina (PAGE_SIZE=50) que a tela tem — o totalizador de
Reconectados, que na tela e recomputado numa segunda passada O(dias)
separada da pagina (`count_reconnected_numbers_for_day` por dia, ver
`_flow_totals_for_period`), aqui e obtido por SOMA INCREMENTAL do valor
`reconectados` de cada linha *enquanto* ela e escrita no arquivo (`_TotalsAccumulator`)
— elimina a segunda passada O(dias) inteira: cada dia paga
`count_reconnected_numbers_for_day` UMA UNICA VEZ (dentro de
`build_indicator_for_day`, para a propria linha), nunca duas. Entregues e
Novos continuam O(1) (`entregues_novos_period_totals`, uma consulta
agregada cada, ja otimizado em REL-13-Q). Equivalencia matematica: a soma
dos valores diarios de Reconectados IGUAL a soma calculada por uma segunda
passada por dia (mesma funcao, mesmo predicado, sem aproximacao) — provada
por comparacao direta em
DailyIndicatorsExportTotalsEquivalenceTests
(config/tests/test_daily_indicators_report.py)."""

import csv
import itertools
import os
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from datetime import timezone as dt_timezone
from typing import Iterator

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.writer.excel import ExcelWriter
from django.http import StreamingHttpResponse
from django.utils import timezone

from core.services.daily_indicators_report_service import (
    entregues_novos_period_totals,
    iter_daily_indicators_report_rows,
)

# Limiar de producao (REL-07/REL-13-C): ate este numero de DIAS exportados
# por arquivo; acima, ZIP com partes numeradas. Testes de fronteira usam o
# parametro `split_threshold` das funcoes abaixo para reduzir o limiar sem
# tocar este valor de producao — nunca reduzido via HTTP.
EXPORT_SPLIT_THRESHOLD = 100_000

COLUMNS = [
    "Data",
    "Pessoas Logadas",
    "% sem Whats",
    "B2B sem Whats",
    "B2C sem Whats",
    "Números Disponíveis",
    "Números Entregues",
    "Reconectados",
    "Novos",
    "Total Descoberto DIA",
]

_CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_ZIP_CONTENT_TYPE = "application/zip"

_BLOCK_SIZE = 65536

_LEADING_WS_CONTROL = re.compile(r"^[\s\x00-\x1f]+")
_FORMULA_TRIGGERS = ("=", "+", "-", "@")

_SLUG = "indicadores_diarios"

_NAO_APLICAVEL_LABEL = "Não aplicável"


@dataclass
class ExportFile:
    path: str
    filename: str
    content_type: str
    temp_dir: str


def _looks_like_formula(value: str) -> bool:
    trimmed = _LEADING_WS_CONTROL.sub("", value)
    return trimmed.startswith(_FORMULA_TRIGGERS)


def _csv_text(value) -> str:
    text = str(value)
    return f"'{text}" if _looks_like_formula(text) else text


def _perc_sem_whats_xlsx_text(value: float) -> str:
    """`92,50%` — decimal com virgula e simbolo % literal (texto, nao
    percentual tipado do Excel): REL-13-C secao 3, "mesma formatacao da
    tela/XLSX", distinta da regra do CSV (numerico puro, sem simbolo, ver
    `_row_to_csv_cells`). Escrito como texto para nao acionar a
    multiplicacao automatica por 100 que o Excel aplica a celulas numericas
    com formato percentual — o valor aqui ja esta em pontos percentuais
    (0-100), nao em fracao (0-1)."""
    return f"{value:.2f}".replace(".", ",") + "%"


def _row_to_csv_cells(row: dict) -> list:
    return [
        row["data"].strftime("%d/%m/%Y"),
        _csv_text(row["pessoas_logadas"]),
        _csv_text(f"{row['perc_sem_whats']:.2f}"),
        _csv_text(row["b2b_sem_whats"]),
        _csv_text(row["b2c_sem_whats"]),
        _csv_text(row["numeros_disponiveis"]),
        _csv_text(row["numeros_entregues"]),
        _csv_text(row["reconectados"]),
        _csv_text(row["novos"]),
        _csv_text(row["total_descoberto_dia"]),
    ]


def _xlsx_date_cell(worksheet, value: date):
    cell = WriteOnlyCell(worksheet, value=value)
    cell.number_format = "DD/MM/YYYY"
    return cell


def _row_to_xlsx_cells(worksheet, row: dict) -> list:
    return [
        _xlsx_date_cell(worksheet, row["data"]),
        row["pessoas_logadas"],
        _perc_sem_whats_xlsx_text(row["perc_sem_whats"]),
        row["b2b_sem_whats"],
        row["b2c_sem_whats"],
        row["numeros_disponiveis"],
        row["numeros_entregues"],
        row["reconectados"],
        row["novos"],
        row["total_descoberto_dia"],
    ]


def _filters_label(filters: dict) -> str:
    return (
        f"Período: {filters['data_inicial'].strftime('%d/%m/%Y')} a "
        f"{filters['data_final'].strftime('%d/%m/%Y')}; "
        f"Incluir excluídos: aceito, sem efeito neste relatório (a "
        f"visibilidade histórica de cada dia já é decidida pelo próprio "
        f"cálculo, sem alternância por checkbox — ver REL-13-C, decisão C)"
    )


class _TotalsAccumulator:
    """Soma `reconectados` de cada linha conforme ela e escrita no arquivo
    (uma unica passada, sem segunda consulta O(dias)) — ver docstring do
    modulo, "Otimizacao aplicada". Envolve o iterator de linhas sem alterar
    seu conteudo nem sua ordem."""

    def __init__(self, rows: Iterator[dict]):
        self._rows = rows
        self.reconectados_total = 0

    def __iter__(self):
        for row in self._rows:
            self.reconectados_total += row["reconectados"]
            yield row


def _write_csv_chunk(path: str, rows: Iterator[dict], threshold: int) -> int:
    """Escreve ate `threshold` DIAS do MESMO iterator (uma unica iteracao
    de todo o periodo dirige toda a exportacao) e retorna quantos foram
    efetivamente escritos. `written < threshold` sinaliza iterator
    esgotado — usado pelo chamador para decidir se ha mais partes, sem
    depender de nenhuma contagem calculada antes da iteracao (mesmo padrao
    corrigido em REL-08-E para evitar divergencia entre contagem previa e
    conjunto realmente iterado)."""
    written = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(COLUMNS)
        for row in itertools.islice(rows, threshold):
            writer.writerow(_row_to_csv_cells(row))
            written += 1
    return written


def _close_abandoned_workbook(workbook) -> None:
    """Fecha um workbook write-only abandonado antes de `save()` — mesmo
    raciocinio e mesma correcao de vazamento de arquivo temporario interno
    do openpyxl (`sheet._writer.out`/`ALL_TEMP_FILES`, ver
    REPORTS_SENIOR_REVIEW achado 2) ja documentados em
    core/services/allocations_export_service."""
    for sheet in workbook.worksheets:
        if getattr(sheet, "_rows", None) is not None and not sheet.closed:
            try:
                sheet.close()
            except Exception:
                pass
        writer = getattr(sheet, "_writer", None)
        if writer is not None:
            try:
                writer.cleanup()
            except Exception:
                pass
    workbook.close()


def _write_xlsx_chunk(path: str, rows: Iterator[dict], threshold: int):
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("Indicadores diários")
    sheet.append(COLUMNS)
    written = 0
    try:
        for row in itertools.islice(rows, threshold):
            sheet.append(_row_to_xlsx_cells(sheet, row))
            written += 1
    except Exception:
        _close_abandoned_workbook(workbook)
        raise
    return workbook, written


def _save_xlsx_workbook(workbook, path: str) -> None:
    """Salva `workbook` em `path` sendo dono explicito do arquivo ZIP de
    saida — nao apenas `workbook.save(path)` — ver docstring completa em
    core/services/allocations_export_service. Fechamento em falha usa
    `archive.close()` normal (metodo publico do `zipfile`), nao mais
    manipulacao direta de `archive.fp`: essa tentativa anterior confundia
    fechar o recurso com entregar o recurso — a entrega HTTP ja e
    impedida pela propagacao da excecao (nunca chega a
    `build_daily_indicators_export_response`) e pela remocao de
    `temp_dir` no `except` externo de `generate_daily_indicators_export`,
    qualquer que seja o conteudo do ZIP parcial no disco. Falha
    SECUNDARIA durante este `close()` de limpeza nao substitui a excecao
    original (por isso e engolida so aqui, ponto estreito e documentado).
    `archive.fp` nao e mais tocado. `workbook.properties.modified` e
    atualizado aqui para preservar o mesmo comportamento de
    `Workbook.save(path)`."""
    workbook.properties.modified = datetime.now(dt_timezone.utc).replace(tzinfo=None)
    archive = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, allowZip64=True)
    try:
        ExcelWriter(workbook, archive).save()
    except Exception:
        try:
            archive.close()
        except Exception:
            pass
        raise


def _finalize_xlsx_chunk(
    workbook,
    path: str,
    *,
    part_index: int,
    num_parts: int,
    total: int,
    count: int,
    generated_at: datetime,
    filters_label: str,
    numeros_entregues_total: int,
    reconectados_total: int,
    novos_total: int,
) -> None:
    context_sheet = workbook.create_sheet("Contexto")
    context_sheet.append(["Campo", "Valor"])
    context_sheet.append(["Relatório", "Indicadores diários"])
    context_sheet.append(["Filtros", filters_label])
    context_sheet.append(["Gerado em", generated_at.strftime("%d/%m/%Y %H:%M:%S")])
    context_sheet.append(["Quantidade total (dias)", total])
    context_sheet.append(["Quantidade nesta parte (dias)", count])
    if num_parts > 1:
        context_sheet.append(["Parte", f"{part_index} de {num_parts}"])
    context_sheet.append(["Números Entregues no período", numeros_entregues_total])
    context_sheet.append(["Reconectados no período", reconectados_total])
    context_sheet.append(["Novos no período", novos_total])
    context_sheet.append(
        [
            "Totalizadores",
            "Somente Números Entregues, Reconectados e Novos são somados no "
            "período — as demais colunas são fotografias de fim de dia, sem "
            f"totalizador ({_NAO_APLICAVEL_LABEL} na tela).",
        ]
    )

    _save_xlsx_workbook(workbook, path)


def _build_base_name(data_inicial: date, data_final: date, generated_at: datetime) -> str:
    return (
        f"{_SLUG}_{data_inicial.strftime('%Y%m%d')}-{data_final.strftime('%Y%m%d')}"
        f"_{generated_at.strftime('%Y%m%dT%H%M%S')}"
    )


def generate_daily_indicators_export(
    *,
    formato: str,
    filters: dict,
    split_threshold: int = EXPORT_SPLIT_THRESHOLD,
) -> ExportFile:
    """Gera o(s) arquivo(s) completo(s) do relatorio Indicadores diarios em
    arquivos temporarios privados, ANTES de qualquer resposta HTTP ser
    iniciada (sem envio parcial). `split_threshold` e a interface tecnica
    explicita de teste para a fronteira de divisao/ZIP (conta DIAS, nao
    registros) — a rota HTTP nunca o passa, mantendo producao em
    EXPORT_SPLIT_THRESHOLD (100.000).

    UMA UNICA iteracao do periodo inteiro (`iter_daily_indicators_report_rows`)
    dirige toda a divisao em partes: nao ha contagem previa separada. Total,
    numero de partes e quantidade por parte sao derivados do que
    efetivamente foi escrito, nunca de uma contagem calculada antes da
    iteracao. O totalizador de Reconectados e acumulado NA MESMA passada
    (`_TotalsAccumulator`, ver docstring do modulo) — nenhuma segunda
    consulta O(dias)."""
    if formato not in _CONTENT_TYPES:
        raise ValueError(f"Formato de exportacao desconhecido: {formato}")

    generated_at = timezone.now()

    numeros_entregues_total, novos_total = entregues_novos_period_totals(
        filters["data_inicial"], filters["data_final"]
    )

    temp_dir = tempfile.mkdtemp(prefix="lineops_daily_indicators_export_")
    xlsx_parts = []  # (workbook, path) — save() adiado ate total/num_parts conhecidos
    try:
        accumulator = _TotalsAccumulator(
            iter_daily_indicators_report_rows(
                data_inicial=filters["data_inicial"],
                data_final=filters["data_final"],
                incluir_excluidos=filters.get("incluir_excluidos", False),
            )
        )
        rows = iter(accumulator)

        base_name = _build_base_name(filters["data_inicial"], filters["data_final"], generated_at)
        filters_label = _filters_label(filters)

        part_paths = []
        part_counts = []
        part_index = 0
        while True:
            part_index += 1
            part_path = os.path.join(temp_dir, f"_part_{part_index:03d}.{formato}")
            if formato == "csv":
                written = _write_csv_chunk(part_path, rows, split_threshold)
            else:
                workbook, written = _write_xlsx_chunk(part_path, rows, split_threshold)

            if written == 0 and part_index > 1:
                # Iterator esgotou exatamente no limiar da parte anterior
                # (ex.: total multiplo exato de split_threshold): esta parte
                # extra e vazia e espuria, nao um resultado real.
                if formato == "csv":
                    os.remove(part_path)
                else:
                    _close_abandoned_workbook(workbook)
                part_index -= 1
                break

            part_paths.append(part_path)
            part_counts.append(written)
            if formato == "xlsx":
                xlsx_parts.append((workbook, part_path))
            if written < split_threshold:
                break

        total = sum(part_counts)
        num_parts = len(part_paths)
        reconectados_total = accumulator.reconectados_total

        if formato == "xlsx":
            for index, (workbook, part_path) in enumerate(xlsx_parts, start=1):
                _finalize_xlsx_chunk(
                    workbook,
                    part_path,
                    part_index=index,
                    num_parts=num_parts,
                    total=total,
                    count=part_counts[index - 1],
                    generated_at=generated_at,
                    filters_label=filters_label,
                    numeros_entregues_total=numeros_entregues_total,
                    reconectados_total=reconectados_total,
                    novos_total=novos_total,
                )

        if num_parts == 1:
            final_path = os.path.join(temp_dir, f"{base_name}.{formato}")
            os.replace(part_paths[0], final_path)
            return ExportFile(
                path=final_path,
                filename=os.path.basename(final_path),
                content_type=_CONTENT_TYPES[formato],
                temp_dir=temp_dir,
            )

        zip_path = os.path.join(temp_dir, f"{base_name}.zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for index, part_path in enumerate(part_paths, start=1):
                arcname = f"{base_name}_parte{index:03d}de{num_parts:03d}.{formato}"
                archive.write(part_path, arcname=arcname)
        for part_path in part_paths:
            os.remove(part_path)

        return ExportFile(
            path=zip_path,
            filename=os.path.basename(zip_path),
            content_type=_ZIP_CONTENT_TYPE,
            temp_dir=temp_dir,
        )
    except Exception:
        # O arquivo ZIP de saida de uma falha durante `_finalize_xlsx_chunk`
        # ja e fechado explicitamente dentro de `_save_xlsx_workbook`
        # (traceback.clear_frames foi tentado antes e descartado por nao
        # alcancar exc.__cause__/exc.__context__ de uma excecao encadeada,
        # ver docstring de `_save_xlsx_workbook` em
        # core/services/allocations_export_service) — resta so fechar os
        # worksheets write-only ainda nao salvos.
        for pending_workbook, _pending_path in xlsx_parts:
            _close_abandoned_workbook(pending_workbook)
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise


def _stream_and_cleanup(path: str, temp_dir: str, block_size: int = _BLOCK_SIZE):
    """Gerador que serve o arquivo ja concluido e remove o diretorio
    temporario ao fim — consumo total, parcial ou desconexao do cliente
    APOS o primeiro chunk. Mesmo raciocinio documentado em
    core/services/pendencies_export_service (o caminho "fechado antes do
    primeiro byte" e coberto por `_CleanupStreamingHttpResponse.close()`,
    nao por este generator)."""
    try:
        with open(path, "rb") as handle:
            while True:
                chunk = handle.read(block_size)
                if not chunk:
                    break
                yield chunk
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


class _CleanupStreamingHttpResponse(StreamingHttpResponse):
    """StreamingHttpResponse que garante a remocao do diretorio temporario
    em `close()`, inclusive quando a resposta e encerrada antes de qualquer
    iteracao. Idempotente via flag propria — mesmo padrao ja usado pelos
    outros cinco exportadores."""

    def __init__(self, streaming_content, *, temp_dir: str, **kwargs):
        super().__init__(streaming_content, **kwargs)
        self._export_temp_dir = temp_dir
        self._export_temp_dir_cleaned = False

    def close(self):
        try:
            super().close()
        finally:
            self._cleanup_export_temp_dir()

    def _cleanup_export_temp_dir(self):
        if self._export_temp_dir_cleaned:
            return
        self._export_temp_dir_cleaned = True
        shutil.rmtree(self._export_temp_dir, ignore_errors=True)


def build_daily_indicators_export_response(
    *,
    formato: str,
    filters: dict,
    split_threshold: int = EXPORT_SPLIT_THRESHOLD,
) -> StreamingHttpResponse:
    export = generate_daily_indicators_export(
        formato=formato, filters=filters, split_threshold=split_threshold
    )
    response = _CleanupStreamingHttpResponse(
        _stream_and_cleanup(export.path, export.temp_dir),
        temp_dir=export.temp_dir,
        content_type=export.content_type,
    )
    response["Content-Disposition"] = f'attachment; filename="{export.filename}"'
    response["Cache-Control"] = "private, no-store"
    return response
