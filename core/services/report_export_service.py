"""Exportacao CSV/XLSX/ZIP do relatorio Selos WhatsApp (REL-08-E).

Fonte da verdade: docs/superpowers/specs/2026-09-22-reports-hub-design.md,
secoes REL-07 (contrato comum) e REL-08-C (contrato especifico). Reaproveita
o motor de consulta de REL-08-Q (core/services/whatsapp_seal_report_service):
mesmo escopo, filtros, ordenacao e projecao de linhas — nenhuma regra e
copiada em outro queryset. Somente leitura: nenhuma escrita de negocio,
model ou migration.
"""

import csv
import itertools
import os
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from datetime import timezone as dt_timezone
from typing import Iterator

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.writer.excel import ExcelWriter
from django.http import StreamingHttpResponse
from django.utils import timezone

from core.services.whatsapp_seal_report_service import (
    WhatsAppSealReportRow,
    iter_whatsapp_seal_report_rows,
)

# Limiar de producao (REL-08-C): ate este numero de linhas de dados por
# arquivo; acima, ZIP com partes numeradas. Testes de fronteira usam o
# parametro `split_threshold` das funcoes abaixo para reduzir o limiar sem
# tocar este valor de producao.
EXPORT_SPLIT_THRESHOLD = 100_000

COLUMNS = [
    "Linha",
    "Tipo do selo",
    "Ativação do selo",
    "Pagamento",
    "E-mail cadastrado",
    "Vencimento",
    "Situação do vencimento",
    "Registro",
]

_DATE_COLUMNS = (False, False, True, False, False, True, False, False)

_CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_ZIP_CONTENT_TYPE = "application/zip"

_BLOCK_SIZE = 65536

_LEADING_WS_CONTROL = re.compile(r"^[\s\x00-\x1f]+")
_FORMULA_TRIGGERS = ("=", "+", "-", "@")

_SLUG = "selos_whatsapp"


@dataclass
class ExportFile:
    path: str
    filename: str
    content_type: str
    temp_dir: str


def _looks_like_formula(value: str) -> bool:
    trimmed = _LEADING_WS_CONTROL.sub("", value)
    return trimmed.startswith(_FORMULA_TRIGGERS)


def _csv_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    return f"'{text}" if _looks_like_formula(text) else text


def _row_to_csv_cells(row: WhatsAppSealReportRow) -> list:
    return [
        _csv_text(row.phone_number),
        _csv_text(row.seal_display),
        row.activated_on.strftime("%d/%m/%Y") if row.activated_on else "",
        _csv_text(row.payment_display),
        _csv_text(row.email),
        row.expires_on.strftime("%d/%m/%Y") if row.expires_on else "",
        _csv_text(row.situation_label),
        _csv_text(row.registro_label),
    ]


def _row_to_xlsx_cells(worksheet, row: WhatsAppSealReportRow) -> list:
    values = [
        row.phone_number,
        row.seal_display,
        row.activated_on,
        row.payment_display or "",
        row.email or "",
        row.expires_on,
        row.situation_label,
        row.registro_label,
    ]
    cells = []
    for value, is_date in zip(values, _DATE_COLUMNS):
        if is_date:
            if value is None:
                cells.append(None)
            else:
                cell = WriteOnlyCell(worksheet, value=value)
                cell.number_format = "DD/MM/YYYY"
                cells.append(cell)
        elif isinstance(value, str) and _looks_like_formula(value):
            # openpyxl infere data_type='f' (formula) para string que comece
            # por '=' atribuida diretamente a uma celula, mesmo em
            # WriteOnlyCell — o contrato (REL-07) exige texto puro, nunca
            # formula, no XLSX. Forcar data_type='s' apos a construcao (nao
            # no `value=` do construtor, que reexecuta a mesma inferencia)
            # grava a string literal sem interpretacao. Mesmo padrao ja
            # validado em core/services/allocations_export_service.py.
            cell = WriteOnlyCell(worksheet, value=value)
            cell.data_type = "s"
            cells.append(cell)
        else:
            cells.append(value)
    return cells


def _filters_label(filters: dict) -> str:
    selo_labels = {"": "Todos", "META_ONE": "META ONE", "META_VERIFIED": "META VERIFICADO"}
    pagamento_labels = {
        "": "Todos",
        "CREDIT_CARD": "Cartão de crédito",
        "GOOGLE_GIFT_CARD": "Gift card Google",
    }
    situacao_labels = {
        "": "Todas",
        "regular": "Vence após 7 dias",
        "warning": "Vence em até 7 dias",
        "today": "Vence hoje",
        "expired": "Vencido",
        "missing_expiry": "Sem vencimento informado",
    }
    return (
        f"Período: {filters['data_inicial'].strftime('%d/%m/%Y')} a "
        f"{filters['data_final'].strftime('%d/%m/%Y')}; "
        f"Tipo de selo: {selo_labels.get(filters.get('selo', ''), filters.get('selo', ''))}; "
        f"Pagamento: {pagamento_labels.get(filters.get('pagamento', ''), filters.get('pagamento', ''))}; "
        f"Situação: {situacao_labels.get(filters.get('situacao', ''), filters.get('situacao', ''))}; "
        f"Linha: {filters.get('linha') or '(vazio)'}; "
        f"Incluir excluídos: {'Sim' if filters.get('incluir_excluidos') else 'Não'}"
    )


def _write_csv_chunk(path: str, rows: Iterator[WhatsAppSealReportRow], threshold: int) -> int:
    """Escreve ate `threshold` linhas do MESMO iterator (uma unica consulta
    para toda a exportacao) e retorna quantas linhas foram efetivamente
    escritas. `written < threshold` sinaliza iterator esgotado — usado pelo
    chamador para decidir se ha mais partes, sem depender de nenhuma
    contagem calculada antes da iteracao (REL-08-E, correcao de risco de
    divergencia entre contagem previa e conjunto realmente iterado)."""
    written = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(COLUMNS)
        for row in itertools.islice(rows, threshold):
            writer.writerow(_row_to_csv_cells(row))
            written += 1
    return written


def _close_abandoned_workbook(workbook) -> None:
    """Fecha um workbook write-only abandonado antes de `save()` — mesmo
    raciocinio e mesma correcao de vazamento de arquivo temporario interno
    do openpyxl (`sheet._writer.out`/`ALL_TEMP_FILES`, ver
    REPORTS_SENIOR_REVIEW achado 2) ja documentados em
    core/services/allocations_export_service."""
    for sheet in workbook.worksheets:
        if getattr(sheet, "_rows", None) is not None and not sheet.closed:
            try:
                sheet.close()
            except Exception:
                pass
        writer = getattr(sheet, "_writer", None)
        if writer is not None:
            try:
                writer.cleanup()
            except Exception:
                pass
    workbook.close()


def _write_xlsx_chunk(path: str, rows: Iterator[WhatsAppSealReportRow], threshold: int):
    """Mesmo contrato de `_write_csv_chunk`, mas para XLSX: os dados sao
    escritos de imediato (write-only, sem materializar linhas em memoria),
    porem a planilha "Contexto" (que depende do total/numero de partes,
    conhecidos somente apos esgotar o iterator inteiro) e adicionada depois
    por `_finalize_xlsx_chunk` — o workbook so e salvo nesse momento."""
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("Selos WhatsApp")
    sheet.append(COLUMNS)
    written = 0
    try:
        for row in itertools.islice(rows, threshold):
            sheet.append(_row_to_xlsx_cells(sheet, row))
            written += 1
    except Exception:
        # Sem isto, uma falha no meio do iterator abandona o worksheet
        # write-only com seu generator interno suspenso; o GC eventual o
        # fecha fora de ordem e emite "ValueError: I/O operation on closed
        # file" como warning nao tratado, alem de vazar o arquivo temporario
        # interno do openpyxl (REPORTS_SENIOR_REVIEW achado 2).
        _close_abandoned_workbook(workbook)
        raise
    return workbook, written


def _save_xlsx_workbook(workbook, path: str) -> None:
    """Salva `workbook` em `path` sendo dono explicito do arquivo ZIP de
    saida — nao apenas `workbook.save(path)` — ver docstring completa em
    core/services/allocations_export_service. Fechamento em falha usa
    `archive.close()` normal (metodo publico do `zipfile`), nao mais
    manipulacao direta de `archive.fp`: essa tentativa anterior confundia
    fechar o recurso com entregar o recurso — a entrega HTTP ja e
    impedida pela propagacao da excecao (nunca chega a
    `build_whatsapp_seal_export_response`) e pela remocao de `temp_dir`
    no `except` externo de `generate_whatsapp_seal_export`, qualquer que
    seja o conteudo do ZIP parcial no disco. Falha SECUNDARIA durante
    este `close()` de limpeza nao substitui a excecao original (por isso
    e engolida so aqui, ponto estreito e documentado). `archive.fp` nao
    e mais tocado. `workbook.properties.modified` e atualizado aqui para
    preservar o mesmo comportamento de `Workbook.save(path)`."""
    workbook.properties.modified = datetime.now(dt_timezone.utc).replace(tzinfo=None)
    archive = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, allowZip64=True)
    try:
        ExcelWriter(workbook, archive).save()
    except Exception:
        try:
            archive.close()
        except Exception:
            pass
        raise


def _finalize_xlsx_chunk(
    workbook,
    path: str,
    *,
    part_index: int,
    num_parts: int,
    total: int,
    count: int,
    reference_date: date,
    generated_at: datetime,
    filters_label: str,
) -> None:
    context_sheet = workbook.create_sheet("Contexto")
    context_sheet.append(["Campo", "Valor"])
    context_sheet.append(["Relatório", "Selos WhatsApp"])
    context_sheet.append(["Filtros", filters_label])
    context_sheet.append(["Referência local", reference_date.strftime("%d/%m/%Y")])
    context_sheet.append(["Gerado em", generated_at.strftime("%d/%m/%Y %H:%M:%S")])
    context_sheet.append(["Quantidade total", total])
    context_sheet.append(["Quantidade nesta parte", count])
    if num_parts > 1:
        context_sheet.append(["Parte", f"{part_index} de {num_parts}"])

    _save_xlsx_workbook(workbook, path)


def _build_base_name(data_inicial: date, data_final: date, generated_at: datetime) -> str:
    return (
        f"{_SLUG}_{data_inicial.strftime('%Y%m%d')}-{data_final.strftime('%Y%m%d')}"
        f"_{generated_at.strftime('%Y%m%dT%H%M%S')}"
    )


def generate_whatsapp_seal_export(
    *,
    formato: str,
    filters: dict,
    split_threshold: int = EXPORT_SPLIT_THRESHOLD,
) -> ExportFile:
    """Gera o(s) arquivo(s) completo(s) do relatorio Selos WhatsApp em
    arquivos temporarios privados, ANTES de qualquer resposta HTTP ser
    iniciada (sem envio parcial). `split_threshold` e a interface tecnica
    explicita de teste para a fronteira de divisao/ZIP — a rota HTTP nunca
    o passa, mantendo producao em EXPORT_SPLIT_THRESHOLD (100.000).

    Uma unica data/hora de referencia e capturada no inicio e usada em toda
    a geracao (filtros, totais e contexto), inclusive se ela atravessar a
    meia-noite durante uma exportacao longa.

    Uma UNICA consulta ao banco (o iterator de `iter_whatsapp_seal_report_rows`)
    dirige toda a divisao em partes: nao ha `count()` previo separado. Total,
    numero de partes e quantidade por parte sao derivados do que
    efetivamente foi escrito, nunca de uma contagem calculada antes da
    iteracao — corrigido apos identificar que duas consultas independentes
    (contagem e iteracao) podiam divergir por mudanca de dados entre elas
    (linha incluida/excluida do conjunto filtrado no intervalo entre as duas
    chamadas), causando arquivo truncado, ZIP incorreto ou falha ao esgotar
    o iterator antes do esperado.
    """
    if formato not in _CONTENT_TYPES:
        raise ValueError(f"Formato de exportacao desconhecido: {formato}")

    reference_date = timezone.localdate()
    generated_at = timezone.now()

    temp_dir = tempfile.mkdtemp(prefix="lineops_reports_export_")
    # Declarado ANTES do try: se uma excecao ocorrer antes/durante o
    # preenchimento desta lista, o except abaixo ainda precisa fecha-la
    # (workbooks write-only ja criados e ainda nao salvos nao podem ficar
    # para o GC fechar fora de ordem, nem vazar o temporario interno do
    # openpyxl — REPORTS_SENIOR_REVIEW achado 2).
    xlsx_parts = []  # (workbook, path) — save() adiado ate total/num_parts conhecidos
    try:
        rows = iter_whatsapp_seal_report_rows(reference_date=reference_date, **filters)

        base_name = _build_base_name(filters["data_inicial"], filters["data_final"], generated_at)
        filters_label = _filters_label(filters)

        part_paths = []
        part_counts = []
        part_index = 0
        while True:
            part_index += 1
            part_path = os.path.join(temp_dir, f"_part_{part_index:03d}.{formato}")
            if formato == "csv":
                written = _write_csv_chunk(part_path, rows, split_threshold)
            else:
                workbook, written = _write_xlsx_chunk(part_path, rows, split_threshold)

            if written == 0 and part_index > 1:
                # Iterator esgotou exatamente no limiar da parte anterior
                # (ex.: total multiplo exato de split_threshold): esta parte
                # extra e vazia e espuria, nao um resultado real. CSV ja
                # grava no disco de imediato (remover o arquivo espurio);
                # XLSX so salva no disco em _finalize_xlsx_chunk (mais
                # abaixo), entao aqui so existe o objeto workbook em
                # memoria (com arquivo temporario interno do openpyxl,
                # nao em temp_dir) — nada em disco (nosso) para remover,
                # mas o workbook write-only precisa ser fechado para nao
                # deixar handle/arquivo interno orfao.
                if formato == "csv":
                    os.remove(part_path)
                else:
                    _close_abandoned_workbook(workbook)
                part_index -= 1
                break

            part_paths.append(part_path)
            part_counts.append(written)
            if formato == "xlsx":
                xlsx_parts.append((workbook, part_path))
            if written < split_threshold:
                break

        total = sum(part_counts)
        num_parts = len(part_paths)

        if formato == "xlsx":
            for index, (workbook, part_path) in enumerate(xlsx_parts, start=1):
                _finalize_xlsx_chunk(
                    workbook,
                    part_path,
                    part_index=index,
                    num_parts=num_parts,
                    total=total,
                    count=part_counts[index - 1],
                    reference_date=reference_date,
                    generated_at=generated_at,
                    filters_label=filters_label,
                )

        if num_parts == 1:
            final_path = os.path.join(temp_dir, f"{base_name}.{formato}")
            os.replace(part_paths[0], final_path)
            return ExportFile(
                path=final_path,
                filename=os.path.basename(final_path),
                content_type=_CONTENT_TYPES[formato],
                temp_dir=temp_dir,
            )

        zip_path = os.path.join(temp_dir, f"{base_name}.zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for index, part_path in enumerate(part_paths, start=1):
                arcname = f"{base_name}_parte{index:03d}de{num_parts:03d}.{formato}"
                archive.write(part_path, arcname=arcname)
        for part_path in part_paths:
            os.remove(part_path)

        return ExportFile(
            path=zip_path,
            filename=os.path.basename(zip_path),
            content_type=_ZIP_CONTENT_TYPE,
            temp_dir=temp_dir,
        )
    except Exception:
        # O arquivo ZIP de saida de uma falha durante `_finalize_xlsx_chunk`
        # ja e fechado explicitamente dentro de `_save_xlsx_workbook`
        # (traceback.clear_frames foi tentado antes e descartado por nao
        # alcancar exc.__cause__/exc.__context__ de uma excecao encadeada,
        # ver docstring de `_save_xlsx_workbook` em
        # core/services/allocations_export_service) — resta so fechar os
        # worksheets write-only ainda nao salvos.
        for pending_workbook, _pending_path in xlsx_parts:
            _close_abandoned_workbook(pending_workbook)
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise


def _stream_and_cleanup(path: str, temp_dir: str, block_size: int = _BLOCK_SIZE):
    """Gerador que serve o arquivo ja concluido e remove o diretorio
    temporario ao fim — consumo total, parcial ou desconexao do cliente
    APOS o primeiro chunk. Um generator so roda seu `finally` ao receber
    `.close()` se ja estiver suspenso em algum `yield`; se `.close()` for
    chamado ANTES de qualquer `next()` (generator no estado GEN_CREATED,
    nunca iniciado), o `finally` NAO executa — comprovado empiricamente
    (`gen.close()` sem `next()` previo nao roda o corpo da funcao). Esse
    caminho (resposta encerrada antes do primeiro byte) e coberto por
    `_CleanupStreamingHttpResponse.close()`, nao por este generator."""
    try:
        with open(path, "rb") as handle:
            while True:
                chunk = handle.read(block_size)
                if not chunk:
                    break
                yield chunk
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


class _CleanupStreamingHttpResponse(StreamingHttpResponse):
    """StreamingHttpResponse que garante a remocao do diretorio temporario
    em `close()`, inclusive quando a resposta e encerrada antes de
    qualquer iteracao (`_stream_and_cleanup` sozinho nao cobre esse caso —
    ver seu docstring). `close()` e metodo publico de HttpResponseBase,
    sempre chamado pelo servidor WSGI real ao final de toda resposta
    (contrato padrao de iteravel fechavel do WSGI), independente de
    streaming_content ter sido consumido. Nao usa nenhum atributo privado
    do Django: apenas estende o metodo publico `close()`. Idempotente via
    flag propria, redundante com `shutil.rmtree(..., ignore_errors=True)`
    caso `close()` seja chamado mais de uma vez."""

    def __init__(self, streaming_content, *, temp_dir: str, **kwargs):
        super().__init__(streaming_content, **kwargs)
        self._export_temp_dir = temp_dir
        self._export_temp_dir_cleaned = False

    def close(self):
        try:
            super().close()
        finally:
            self._cleanup_export_temp_dir()

    def _cleanup_export_temp_dir(self):
        if self._export_temp_dir_cleaned:
            return
        self._export_temp_dir_cleaned = True
        shutil.rmtree(self._export_temp_dir, ignore_errors=True)


def build_whatsapp_seal_export_response(
    *,
    formato: str,
    filters: dict,
    split_threshold: int = EXPORT_SPLIT_THRESHOLD,
) -> StreamingHttpResponse:
    export = generate_whatsapp_seal_export(
        formato=formato, filters=filters, split_threshold=split_threshold
    )
    response = _CleanupStreamingHttpResponse(
        _stream_and_cleanup(export.path, export.temp_dir),
        temp_dir=export.temp_dir,
        content_type=export.content_type,
    )
    response["Content-Disposition"] = f'attachment; filename="{export.filename}"'
    response["Cache-Control"] = "private, no-store"
    return response
